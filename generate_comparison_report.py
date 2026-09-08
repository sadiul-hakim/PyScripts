import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OLD_CSV_PATH = r"C:\Users\Sadiul Hakim\OneDrive\Desktop\Data\Old Data.csv"
NEW_CSV_PATH = r"C:\Users\Sadiul Hakim\OneDrive\Desktop\Data\New Data_Filtered.csv"
ORIG_EXCEL_TEMPLATE = r"C:\Users\Sadiul Hakim\OneDrive\Desktop\Data\QR Billing Details Report with Data (New System) Apr 2026.xlsx"

EXCEL_OUTPUT_PATH = r"C:\Users\Sadiul Hakim\OneDrive\Desktop\Data\QR_Billing_Comparison_Discrepancies.xlsx"
DESKTOP_EXCEL_PATH = r"C:\Users\Sadiul Hakim\OneDrive\Desktop\QR_Billing_Comparison_Discrepancies.xlsx"


def get_original_headers_and_merges():
    """Extract full 5-tier headers and merged ranges from original template."""
    wb_orig = openpyxl.load_workbook(ORIG_EXCEL_TEMPLATE, data_only=True)
    ws_orig = wb_orig.active

    # Extract 5 header rows (1-based index)
    header_matrix = []
    for r in range(1, 6):
        row_vals = [ws_orig.cell(r, c).value for c in range(1, 321)]
        header_matrix.append(row_vals)

    # Extract merged ranges in rows 1-5
    merged_ranges = []
    for rng in ws_orig.merged_cells.ranges:
        if rng.min_row <= 5:
            merged_ranges.append((rng.min_row, rng.min_col, min(rng.max_row, 5), rng.max_col))

    wb_orig.close()
    return header_matrix, merged_ranges


def build_hierarchical_column_names(header_matrix):
    """Build clear composite names for the summary tab."""
    num_cols = 320
    col_names = []
    for c_idx in range(num_cols):
        if c_idx == 0:
            col_names.append("Date")
        elif c_idx == 1:
            col_names.append("Deal")
        else:
            parts = []
            for r_idx in range(5):
                v = header_matrix[r_idx][c_idx]
                if v is not None:
                    v_str = str(v).strip()
                    if v_str and v_str != "-" and v_str != "QR Billing Details Report" and v_str not in parts:
                        parts.append(v_str)
            col_names.append(" - ".join(parts) if parts else f"Col_{c_idx}")
    return col_names


def generate_excel_report():
    print(f"Reading Old Data: {OLD_CSV_PATH}")
    with open(OLD_CSV_PATH, "r", encoding="utf-8-sig", errors="ignore") as f:
        old_rows = list(csv.reader(f))

    print(f"Reading New Data: {NEW_CSV_PATH}")
    with open(NEW_CSV_PATH, "r", encoding="utf-8-sig", errors="ignore") as f:
        new_rows = list(csv.reader(f))

    print("Extracting full 5-tier header structure and merges...")
    header_matrix, orig_merges = get_original_headers_and_merges()
    col_titles = build_hierarchical_column_names(header_matrix)

    old_data = {r[1].strip(): r for r in old_rows[5:] if len(r) > 1 and r[1].strip()}
    new_data = {r[1].strip(): r for r in new_rows[5:] if len(r) > 1 and r[1].strip()}
    ordered_deals = [r[1].strip() for r in new_rows[5:] if len(r) > 1 and r[1].strip()]

    num_metrics = 320 - 2  # 318 metric columns

    # -------------------------------------------------------------
    # Calculate Column Summary
    # -------------------------------------------------------------
    print("Calculating summary statistics per column...")
    summary_rows = []
    for c in range(2, 320):
        tot_old = 0.0
        tot_new = 0.0
        diff_count = 0
        zero_in_new = 0
        zero_in_old = 0
        mismatch_count = 0

        for deal in ordered_deals:
            r_old = old_data.get(deal, [])
            r_new = new_data.get(deal, [])
            vo = r_old[c].strip() if c < len(r_old) else ""
            vn = r_new[c].strip() if c < len(r_new) else ""

            try:
                fo = float(vo) if vo != "" else 0.0
                fn = float(vn) if vn != "" else 0.0
                tot_old += fo
                tot_new += fn
                if abs(fo - fn) > 1e-4:
                    diff_count += 1
                    if fn == 0.0 and fo != 0.0:
                        zero_in_new += 1
                    elif fo == 0.0 and fn != 0.0:
                        zero_in_old += 1
                    else:
                        mismatch_count += 1
            except ValueError:
                if vo != vn:
                    diff_count += 1
                    mismatch_count += 1

        gap_tot = tot_old - tot_new
        summary_rows.append([
            c,
            col_titles[c],
            diff_count,
            zero_in_new,
            zero_in_old,
            mismatch_count,
            round(tot_old, 2),
            round(tot_new, 2),
            round(gap_tot, 2)
        ])

    # -------------------------------------------------------------
    # GENERATE FULL STRUCTURED EXCEL WORKBOOK (.XLSX)
    # -------------------------------------------------------------
    print(f"\nBuilding Formatted Excel Workbook with full 5-tier headers: {EXCEL_OUTPUT_PATH} ...")
    wb = openpyxl.Workbook()

    # Styling Palette
    hdr1_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hdr1_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    hdr2_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    hdr2_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    hdr3_fill = PatternFill(start_color="8EA9DB", end_color="8EA9DB", fill_type="solid")
    hdr3_font = Font(name="Calibri", size=10, bold=True, color="000000")

    hdr4_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    hdr4_font = Font(name="Calibri", size=10, bold=True, color="000000")

    hdr5_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    hdr5_font = Font(name="Calibri", size=9, bold=True, color="000000")

    sub_old_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    sub_new_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    sub_gap_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    sub_font = Font(name="Calibri", size=9, bold=True, color="000000")

    diff_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    diff_font = Font(name="Calibri", size=9, bold=True, color="9C0006")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    ws_comp = wb.active
    ws_comp.title = "Comparison_Full_Hierarchy"

    # Populate cell values for Rows 1 to 5
    for r_idx in range(5):
        row_num = r_idx + 1
        ws_comp.cell(row=row_num, column=1, value=header_matrix[r_idx][0])
        ws_comp.cell(row=row_num, column=2, value=header_matrix[r_idx][1])
        for c_idx in range(2, 320):
            orig_col = c_idx + 1  # 1-based original column
            new_start_col = 3 + (orig_col - 3) * 3
            val = header_matrix[r_idx][c_idx]
            if val is not None and str(val).strip() != "":
                ws_comp.cell(row=row_num, column=new_start_col, value=val)

    # Apply merged ranges from original file mapped to 3x expansion
    print("Mapping and applying merged header ranges...")
    mapped_merges = set()

    for min_r, min_c, max_r, max_c in orig_merges:
        if min_c == 1 and max_c == 1:
            mapped_merges.add((1, 1, 6, 1))
        elif min_c == 2 and max_c == 2:
            mapped_merges.add((1, 2, 6, 2))
        elif min_c >= 3:
            new_min_c = 3 + (min_c - 3) * 3
            new_max_c = 3 + (max_c - 3) * 3 + 2
            mapped_merges.add((min_r, new_min_c, max_r, new_max_c))

    # For cells in rows 1 to 5 that have values and are not yet merged across their 3 columns
    for r_idx in range(5):
        row_num = r_idx + 1
        for orig_col in range(3, 321):
            new_start_col = 3 + (orig_col - 3) * 3
            new_end_col = new_start_col + 2
            val = header_matrix[r_idx][orig_col - 1]
            if val is not None and str(val).strip() != "" and str(val).strip() != "-":
                covered = False
                for mr1, mc1, mr2, mc2 in mapped_merges:
                    if mr1 <= row_num <= mr2 and mc1 <= new_start_col and new_end_col <= mc2:
                        covered = True
                        break
                if not covered:
                    mapped_merges.add((row_num, new_start_col, row_num, new_end_col))

    # Apply all merges
    for min_r, min_c, max_r, max_c in mapped_merges:
        try:
            ws_comp.merge_cells(start_row=min_r, start_column=min_c, end_row=max_r, end_column=max_c)
        except Exception:
            pass

    # Style Header Rows 1 to 5
    tier_fills = [hdr1_fill, hdr2_fill, hdr3_fill, hdr4_fill, hdr5_fill]
    tier_fonts = [hdr1_font, hdr2_font, hdr3_font, hdr4_font, hdr5_font]

    for r_idx in range(5):
        row_num = r_idx + 1
        f_fill = tier_fills[r_idx]
        f_font = tier_fonts[r_idx]
        ws_comp.row_dimensions[row_num].height = 24
        for col_idx in range(1, 3 + num_metrics * 3):
            cell = ws_comp.cell(row=row_num, column=col_idx)
            cell.fill = f_fill
            cell.font = f_font
            cell.alignment = center_align
            cell.border = thin_border

    # Format Date & Deal Header Cells (Top-Left of merged range A1:A6 and B1:B6)
    c_date = ws_comp.cell(row=1, column=1, value="Date")
    c_date.fill = hdr1_fill
    c_date.font = hdr1_font
    c_date.alignment = center_align

    c_deal = ws_comp.cell(row=1, column=2, value="Deal")
    c_deal.fill = hdr1_fill
    c_deal.font = hdr1_font
    c_deal.alignment = center_align

    # Row 6: Sub-headers (Old, New, Gap)
    ws_comp.row_dimensions[6].height = 20

    for i in range(num_metrics):
        col_start = 3 + i * 3
        c_old = ws_comp.cell(row=6, column=col_start, value="Old")
        c_old.fill = sub_old_fill
        c_old.font = sub_font
        c_old.alignment = center_align
        c_old.border = thin_border

        c_new = ws_comp.cell(row=6, column=col_start + 1, value="New")
        c_new.fill = sub_new_fill
        c_new.font = sub_font
        c_new.alignment = center_align
        c_new.border = thin_border

        c_gap = ws_comp.cell(row=6, column=col_start + 2, value="Gap")
        c_gap.fill = sub_gap_fill
        c_gap.font = sub_font
        c_gap.alignment = center_align
        c_gap.border = thin_border

    # Freeze panes at C7 so all 6 header rows stay visible
    ws_comp.freeze_panes = "C7"

    # Write Data Rows (Row 7+)
    print("Writing comparison data rows to Excel...")
    row_num = 7
    for deal in ordered_deals:
        r_old = old_data.get(deal, [])
        r_new = new_data.get(deal, [])

        date_val = r_new[0].strip() if len(r_new) > 0 and r_new[0].strip() else (r_old[0].strip() if len(r_old) > 0 else "")
        ws_comp.cell(row=row_num, column=1, value=date_val).alignment = center_align
        ws_comp.cell(row=row_num, column=2, value=deal).alignment = left_align
        ws_comp.cell(row=row_num, column=1).border = thin_border
        ws_comp.cell(row=row_num, column=2).border = thin_border

        for i, c in enumerate(range(2, 320)):
            col_start = 3 + i * 3
            vo = r_old[c].strip() if c < len(r_old) else ""
            vn = r_new[c].strip() if c < len(r_new) else ""

            is_diff = False
            try:
                fo = float(vo) if vo != "" else 0.0
                fn = float(vn) if vn != "" else 0.0
                gap = fo - fn
                if abs(gap) < 1e-6:
                    gap = 0.0
                else:
                    is_diff = True

                c_o = ws_comp.cell(row=row_num, column=col_start, value=fo if vo != "" else 0)
                c_n = ws_comp.cell(row=row_num, column=col_start + 1, value=fn if vn != "" else 0)
                c_g = ws_comp.cell(row=row_num, column=col_start + 2, value=gap)
                c_o.alignment = right_align
                c_n.alignment = right_align
                c_g.alignment = right_align
            except ValueError:
                is_diff = (vo != vn)
                c_o = ws_comp.cell(row=row_num, column=col_start, value=vo)
                c_n = ws_comp.cell(row=row_num, column=col_start + 1, value=vn)
                c_g = ws_comp.cell(row=row_num, column=col_start + 2, value="" if not is_diff else "MISMATCH")

            c_o.border = thin_border
            c_n.border = thin_border
            c_g.border = thin_border

            if is_diff:
                c_g.fill = diff_fill
                c_g.font = diff_font

        row_num += 1

    # -------------------------------------------------------------
    # Sheet 2: Summary by Column
    # -------------------------------------------------------------
    print("Creating Summary_By_Column sheet...")
    ws_sum = wb.create_sheet(title="Summary_By_Column")
    sum_headers = [
        "Column Index",
        "Full Billing Hierarchy Path",
        "Discrepancy Count",
        "Zero in New (Missing)",
        "Zero in Old",
        "Mismatch / Rounding",
        "Old Total Sum",
        "New Total Sum",
        "Net Gap (Old - New)"
    ]
    ws_sum.append(sum_headers)
    for col_idx in range(1, len(sum_headers) + 1):
        c_cell = ws_sum.cell(row=1, column=col_idx)
        c_cell.fill = hdr1_fill
        c_cell.font = hdr1_font
        c_cell.alignment = center_align

    for s_row in summary_rows:
        ws_sum.append(s_row)

    for r in range(2, len(summary_rows) + 2):
        ws_sum.cell(row=r, column=1).alignment = center_align
        ws_sum.cell(row=r, column=3).alignment = right_align
        ws_sum.cell(row=r, column=4).alignment = right_align
        ws_sum.cell(row=r, column=5).alignment = right_align
        ws_sum.cell(row=r, column=6).alignment = right_align
        for col_f in range(7, 10):
            cell = ws_sum.cell(row=r, column=col_f)
            cell.number_format = "#,##0.00"
            cell.alignment = right_align
            if col_f == 9 and abs(cell.value or 0) > 1:
                cell.font = Font(name="Calibri", size=10, bold=True, color="9C0006")
                cell.fill = diff_fill

    ws_sum.column_dimensions["A"].width = 15
    ws_sum.column_dimensions["B"].width = 65
    ws_sum.column_dimensions["C"].width = 20
    ws_sum.column_dimensions["D"].width = 22
    ws_sum.column_dimensions["E"].width = 20
    ws_sum.column_dimensions["F"].width = 24
    ws_sum.column_dimensions["G"].width = 22
    ws_sum.column_dimensions["H"].width = 22
    ws_sum.column_dimensions["I"].width = 24
    ws_sum.freeze_panes = "A2"

    for save_path in [EXCEL_OUTPUT_PATH, DESKTOP_EXCEL_PATH]:
        try:
            print(f"Saving Excel file to: {save_path} ...")
            wb.save(save_path)
            print(f"[OK] Saved: {save_path}")
        except PermissionError:
            alt_path = save_path.replace(".xlsx", "_Updated.xlsx")
            print(f"[NOTE] '{save_path}' is currently open in Excel. Saving to: {alt_path}")
            wb.save(alt_path)
            print(f"[OK] Saved: {alt_path}")


if __name__ == "__main__":
    generate_excel_report()
